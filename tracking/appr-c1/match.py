# -*- coding: utf-8 -*-
"""Cruce nombre CONFIG <-> fila del Excel, por nombre normalizado."""
import json, os, re, sys, unicodedata
import openpyxl

XL = r'C:\www\code-corhuila\repos_2026-b\activity-management\list-data\calification-term-1.xlsx'
HERE = os.path.dirname(os.path.abspath(__file__))
rows = json.load(open(os.path.join(HERE, "appr-resultados.json"), encoding="utf-8"))

def norm(s):
    s = unicodedata.normalize("NFKD", str(s))
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r'[^a-z ]', ' ', s.lower())
    return " ".join(s.split())

wb = openpyxl.load_workbook(XL)
ws = wb["P-M"]

# 1) verificar encabezado EN VIVO
hdr = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
col_appr = hdr.get("appr")
col_nom = hdr.get("nombre")
print(f"Encabezado vivo: columna 'appr' -> {openpyxl.utils.get_column_letter(col_appr)} "
      f"| 'nombre' -> {openpyxl.utils.get_column_letter(col_nom)}")
if openpyxl.utils.get_column_letter(col_appr) != "P":
    print("!! OJO: 'appr' YA NO ESTA EN P. Se usara la columna real detectada.")

excel = {}
for r in range(2, ws.max_row + 1):
    v = ws.cell(r, col_nom).value
    if v:
        excel.setdefault(norm(v), []).append(r)

# nombre del CONFIG -> tokens; el Excel usa APELLIDOS NOMBRES, el CONFIG Nombres Apellidos
def tokens(s):
    return frozenset(norm(s).split())

excel_tok = {}
for r in range(2, ws.max_row + 1):
    v = ws.cell(r, col_nom).value
    if v:
        excel_tok.setdefault(tokens(v), []).append(r)

out, problemas = [], []
usadas = set()
for rec in rows:
    name = rec["name"]
    if not name:
        problemas.append((rec["user"], "sin CONFIG -> sin nombre para cruzar"))
        out.append({**rec, "row": None, "excel_name": None, "match": "SIN CONFIG"})
        continue
    t = tokens(name)
    cands = excel_tok.get(t, [])
    how = "exacto (mismos tokens)"
    if not cands:  # subconjunto / superconjunto
        cands = [r for tt, rr in excel_tok.items() if t <= tt or tt <= t for r in rr]
        how = "por subconjunto de tokens"
    if len(cands) != 1:
        problemas.append((rec["user"], f"{name!r} cruza con {len(cands)} filas {cands}"))
        out.append({**rec, "row": None, "excel_name": None, "match": f"AMBIGUO ({len(cands)})"})
        continue
    r = cands[0]
    if r in usadas:
        problemas.append((rec["user"], f"fila {r} ya asignada a otro estudiante"))
        out.append({**rec, "row": None, "excel_name": None, "match": "FILA DUPLICADA"})
        continue
    usadas.add(r)
    out.append({**rec, "row": r, "excel_name": str(ws.cell(r, col_nom).value).strip(),
                "match": how, "actual": ws.cell(r, col_appr).value})

print(f"\n{'FILA':>5}  {'NOMBRE EN EL EXCEL':38} {'CONFIG':30} {'actual':>7} {'nueva':>7}  cruce")
print("-" * 110)
for o in sorted(out, key=lambda x: (x["row"] is None, x["row"] or 0)):
    print(f"{str(o['row'] or '--'):>5}  {(o['excel_name'] or '-')[:37]:38} "
          f"{(o['name'] or '['+o['user']+']')[:29]:30} {str(o.get('actual') or '-'):>7} "
          f"{o['appr']:>7.2f}  {o['match']}")

libres = [r for r in range(2, ws.max_row + 1) if r not in usadas and ws.cell(r, col_nom).value]
print(f"\nFilas del Excel SIN estudiante cruzado: {libres}")
for r in libres:
    print(f"   fila {r}: {ws.cell(r, col_nom).value!r}  (appr actual: {ws.cell(r, col_appr).value!r})")
print(f"\nProblemas: {len(problemas)}")
for p in problemas:
    print("   ", p)

json.dump(out, open(os.path.join(HERE, "appr-cruce.json"), "w", encoding="utf-8"),
          ensure_ascii=False, indent=1)
