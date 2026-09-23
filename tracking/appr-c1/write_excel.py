# -*- coding: utf-8 -*-
"""Escribe la appr en la hoja P-M, columna 'appr'. Por defecto SIMULA.

    py -3.11 write_excel.py           -> simulacion (no toca el archivo)
    py -3.11 write_excel.py --write   -> escribe de verdad

Reglas: solo toca la columna 'appr' (localizada por encabezado vivo, no por
posicion); nunca sobrescribe un valor existente que difiera del calculado;
no escribe filas sin cruce univoco.
"""
import json, os, shutil, sys, datetime as dt
import openpyxl

XL = r'C:\www\code-corhuila\repos_2026-b\activity-management\list-data\calification-term-1.xlsx'
HERE = os.path.dirname(os.path.abspath(__file__))
WRITE = "--write" in sys.argv

rows = json.load(open(os.path.join(HERE, "appr-cruce.json"), encoding="utf-8"))
wb = openpyxl.load_workbook(XL)
ws = wb["P-M"]

hdr = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
if "appr" not in hdr:
    sys.exit("ABORTA: no existe una columna 'appr' en la fila de encabezado de P-M.")
col = hdr["appr"]
letra = openpyxl.utils.get_column_letter(col)
print(f"Columna 'appr' detectada en vivo: {letra}{1}  (valor del encabezado: {ws.cell(1,col).value!r})")
if letra != "P":
    print(f"!! AVISO: 'appr' se movio de P a {letra}. Se escribe en {letra}.")

plan, saltos = [], []
for r in sorted(rows, key=lambda x: x["row"] or 999):
    if not r["row"]:
        saltos.append((r["name"] or r["user"], f"sin cruce univoco ({r['match']})", r["appr"]))
        continue
    cell = ws.cell(r["row"], col)
    actual = cell.value
    if actual is not None and str(actual).strip() != "":
        if abs(float(actual) - r["appr"]) < 1e-9:
            saltos.append((r["excel_name"], f"ya tiene {actual} = calculado, no se toca", r["appr"]))
        else:
            saltos.append((r["excel_name"], f"YA TIENE {actual} != calculado {r['appr']} -> NO se sobrescribe", r["appr"]))
        continue
    plan.append((r["row"], f"{letra}{r['row']}", r["excel_name"], actual, r["appr"]))

print(f"\n=== SE ESCRIBIRIAN {len(plan)} CELDAS ===")
print(f"{'CELDA':>7}  {'ESTUDIANTE':40} {'antes':>7} {'->':^4} {'appr':>6}")
print("-" * 72)
for _, ref, name, antes, val in plan:
    print(f"{ref:>7}  {name[:39]:40} {str(antes or 'vacia'):>7} {'->':^4} {val:>6.2f}")

print(f"\n=== NO SE TOCAN ({len(saltos)}) ===")
for name, why, val in saltos:
    print(f"  {name[:42]:44} calc={val:.2f}  |  {why}")

if not WRITE:
    print("\n[SIMULACION] No se modifico el archivo. Para escribir: --write")
    sys.exit(0)

bak = XL.replace(".xlsx", f".bak-{dt.datetime.now():%Y%m%d-%H%M%S}.xlsx")
shutil.copy2(XL, bak)
print(f"\nRespaldo: {bak}")
for row, ref, name, antes, val in plan:
    ws.cell(row, col).value = val
try:
    wb.save(XL)
except PermissionError:
    sys.exit("ABORTA: el archivo esta abierto en Excel. Cierralo y vuelve a correr.")
print(f"Escritas {len(plan)} celdas en {XL} (hoja P-M, columna {letra}).")

chk = openpyxl.load_workbook(XL)["P-M"]
malas = [ref for row, ref, n, a, v in plan if abs(float(chk.cell(row, col).value) - v) > 1e-9]
print("Verificacion tras guardar:", "OK, todas las celdas coinciden" if not malas else f"FALLAN {malas}")
