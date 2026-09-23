"""Convierte cada entrega guardada (pdf/docx/md/...) a texto plano en text/."""
import io, json, os, sys, zipfile, re

HERE = os.path.dirname(os.path.abspath(__file__))
FILES, TEXT = os.path.join(HERE, "files"), os.path.join(HERE, "text")
os.makedirs(TEXT, exist_ok=True)
report = []

def pdf_text(p):
    try:
        import fitz
        d = fitz.open(p)
        return "\n".join(pg.get_text() for pg in d)
    except Exception as e:
        try:
            import pdfplumber
            with pdfplumber.open(p) as d:
                return "\n".join((pg.extract_text() or "") for pg in d.pages)
        except Exception as e2:
            return f"[[NO EXTRAIBLE: {e} / {e2}]]"

def docx_text(p):
    try:
        import docx
        d = docx.Document(p)
        parts = [x.text for x in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return "\n".join(parts)
    except Exception as e:
        return f"[[NO EXTRAIBLE: {e}]]"

for user in sorted(os.listdir(FILES)):
    d = os.path.join(FILES, user)
    if not os.path.isdir(d):
        continue
    os.makedirs(os.path.join(TEXT, user), exist_ok=True)
    for fn in sorted(os.listdir(d)):
        src = os.path.join(d, fn)
        low = fn.lower()
        if low.endswith(".pdf"):
            t = pdf_text(src)
        elif low.endswith(".docx"):
            t = docx_text(src)
        elif low.endswith((".png", ".jpg", ".jpeg", ".zip", ".gif", ".webp")):
            t = f"[[BINARIO {os.path.getsize(src)}B — no texto]]"
        else:
            try:
                t = open(src, encoding="utf-8", errors="replace").read()
            except Exception as e:
                t = f"[[NO LEIBLE: {e}]]"
        out = os.path.join(TEXT, user, fn + ".txt")
        with open(out, "w", encoding="utf-8") as fh:
            fh.write(t)
        words = len(re.findall(r"\w+", t)) if not t.startswith("[[") else 0
        report.append({"user": user, "file": fn, "chars": len(t), "words": words,
                       "ok": not t.startswith("[[")})

with open(os.path.join(HERE, "raw", "text_index.json"), "w", encoding="utf-8") as f:
    json.dump(report, f, ensure_ascii=False, indent=1)

bad = [r for r in report if not r["ok"]]
print(f"{len(report)} archivos -> texto. {len(bad)} sin texto extraible (binarios/imagenes).")
for r in report:
    if r["ok"] and r["words"] < 40:
        print(f"  MUY CORTO ({r['words']}w): {r['user']}/{r['file']}")
